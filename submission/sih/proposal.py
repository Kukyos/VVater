"""Fill the SIH 2026 proposal template from the repository.

    python -m server.eval.run_eval --json          # first, so the numbers are current
    python submission/sih/proposal.py              # figures + final/VVater-SIH2026-proposal.docx
    python submission/sih/proposal.py --check      # fails on leftover template text or word limits
    powershell -ExecutionPolicy Bypass -File submission/sih/topdf.ps1

`SIH Project Proposal Template.docx` must sit in the repo root, like the deck's template.
The template is opened and filled in place, so its fonts, header table and section
headings are the portal's own. The gray instruction runs are removed by colour, not by
matching their text, and the empty writing space under each heading goes with them.

Every number is read from `data/eval-latest.json` (hard rule 1); a missing key is a crash.
The three diagrams are the deck's own (`figures.flowchart`, `architecture`, `band_chart`),
rendered at portrait width.
"""

from __future__ import annotations

import copy
import json
import subprocess
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from docx.table import Table
from docx.text.paragraph import Paragraph

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]
FIGURES = HERE / "figures"
TEMPLATE = ROOT / "SIH Project Proposal Template.docx"
OUT = HERE / "final" / "VVater-SIH2026-proposal.docx"
EVAL = ROOT / "data" / "eval-latest.json"
HOSTING = ROOT / "data" / "hosting-latest.json"   # the request list, for the denominator

sys.path.insert(0, str(HERE))
import figures  # noqa: E402  the deck's diagrams, reused
from build import TITLE_FIELDS  # noqa: E402  same ID, title and team as the deck

GRAY = "A6A6A6"
BODY_PT = 11
AUTHOR = "A Mohamed Armaan"
REPO = "github.com/Kukyos/VVater"


def n(v, d=2) -> str:
    return f"{v:,.{d}f}"


# ------------------------------------------------------------------ figures

def diagrams(e: dict) -> None:
    """The deck's three diagrams as standalone pages, rendered by render.py. deck.css is
    inlined rather than linked: render.py sizes a page that links it as a whole slide."""
    css = (FIGURES / "deck.css").read_text(encoding="utf-8")
    css += "\nhtml, body { width: auto; height: auto; overflow: visible; background: #fff; }\n"
    pages = {
        "doc-flow": figures.flowchart(e["v2"]),
        "doc-arch": figures.architecture(e["v2"]["hosting"]),
        "doc-bands": figures.band_chart(e["depth_bands"], width=1600, height=760),
    }
    for name, svg in pages.items():
        (FIGURES / f"{name}.html").write_text(
            '<!doctype html><html><head><meta charset="utf-8">'
            '<link href="https://fonts.googleapis.com/css2?family=IBM+Plex+Sans:wght@400;600'
            '&family=IBM+Plex+Mono&display=swap" rel="stylesheet">'
            f"<style>{css}</style></head><body>{svg}</body></html>", encoding="utf-8")
    subprocess.run([sys.executable, str(HERE / "render.py"), *pages], check=True)


# ------------------------------------------------------------------ docx helpers

def new_para_after(anchor, text: str = "", label: str | None = None, *, size=BODY_PT,
                   bold=False, italic=False, align=None, indent_cm=None) -> Paragraph:
    """A body paragraph after `anchor` (a Paragraph or Table). `label` is a bold lead-in."""
    el = anchor._p if isinstance(anchor, Paragraph) else anchor._tbl
    p_el = el.makeelement(qn("w:p"), {})
    el.addnext(p_el)
    p = Paragraph(p_el, anchor._parent)
    p.paragraph_format.space_after = Pt(6)
    if align is not None:
        p.alignment = align
    if indent_cm is not None:
        p.paragraph_format.left_indent = Cm(indent_cm)
        p.paragraph_format.first_line_indent = Cm(-0.45)
    if label:
        r = p.add_run(label + " ")
        r.bold, r.font.size = True, Pt(size)
    if text:
        r = p.add_run(text)
        r.bold, r.italic, r.font.size = bold, italic, Pt(size)
    return p


class Writer:
    """Appends blocks in order after an anchor."""

    def __init__(self, anchor):
        self.at = anchor

    def p(self, text, label=None, **kw):
        self.at = new_para_after(self.at, text, label, **kw)
        return self

    def bullets(self, items):
        for item in items:
            label, text = item if isinstance(item, tuple) else (None, item)
            self.at = new_para_after(self.at, "", None, indent_cm=0.9)
            self.at.paragraph_format.space_after = Pt(3)
            self.at.add_run("•  ").font.size = Pt(BODY_PT)
            if label:
                r = self.at.add_run(label + " ")
                r.bold, r.font.size = True, Pt(BODY_PT)
            self.at.add_run(text).font.size = Pt(BODY_PT)
        return self

    def figure(self, image: Path, caption: str, width_cm=16.4):
        self.at = new_para_after(self.at, align=WD_ALIGN_PARAGRAPH.CENTER)
        self.at.paragraph_format.keep_with_next = True
        self.at.add_run().add_picture(str(image), width=Cm(width_cm))
        self.at = new_para_after(self.at, caption, size=9, italic=True,
                                 align=WD_ALIGN_PARAGRAPH.CENTER)
        return self


def strip_gray(doc) -> None:
    """Remove every gray instruction run; a paragraph left empty by that goes too."""
    for p in list(doc.paragraphs):
        gray = [r for r in p.runs if r.font.color is not None and r.font.color.type
                and str(r.font.color.rgb) == GRAY]
        if not gray:
            continue
        for r in gray:
            r._r.getparent().remove(r._r)
        if not p.text.strip():
            p._p.getparent().remove(p._p)


def drop_empty_paragraphs(doc) -> None:
    """The template's blank writing space. Keeps the title (it ends in a break) and any
    paragraph holding a picture or section properties."""
    for p in list(doc.paragraphs):
        el = p._p
        if p.text.strip() or el.xpath(".//w:drawing") or el.xpath(".//w:sectPr"):
            continue
        el.getparent().remove(el)


def find(doc, starts: str) -> Paragraph:
    hits = [p for p in doc.paragraphs if p.text.strip().startswith(starts)]
    if len(hits) != 1:
        raise SystemExit(f"template anchor {starts!r} matched {len(hits)} paragraphs")
    return hits[0]


def set_text(p: Paragraph, text: str, bold=None) -> None:
    """Replace a paragraph's text, keeping its first run's formatting."""
    runs = p.runs
    for r in runs[1:]:
        r._r.getparent().remove(r._r)
    if runs:
        runs[0].text = text
        if bold is not None:
            runs[0].bold = bold
    else:
        p.add_run(text).bold = bold


def set_cell(cell, text: str, bold=False, size=10) -> None:
    for extra in cell.paragraphs[1:]:
        extra._p.getparent().remove(extra._p)
    p = cell.paragraphs[0]
    for r in p.runs:
        r._r.getparent().remove(r._r)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run(text)
    r.bold, r.font.size = bold, Pt(size)


def keep_rows_whole(table: Table) -> None:
    """No row breaks across a page, and the header row repeats on the next one."""
    for i, row in enumerate(table.rows):
        trPr = row._tr.get_or_add_trPr()
        trPr.append(trPr.makeelement(qn("w:cantSplit"), {}))
        if i == 0:
            trPr.append(trPr.makeelement(qn("w:tblHeader"), {}))


def fill_table(table: Table, rows: list[list[str]], first_bold=True, footer=0) -> None:
    """Fill the body rows (row 0 is the header; the last `footer` rows are left alone),
    adding or removing rows to fit by copying the first body row's XML."""
    end = len(table.rows) - footer
    body = [r._tr for r in table.rows[1:end]]
    while len(body) < len(rows):
        body[-1].addnext(copy.deepcopy(body[0]))
        body.append(body[-1].getnext())
    for extra in body[len(rows):]:
        extra.getparent().remove(extra)
    for tr, values in zip(table.rows[1:1 + len(rows)], rows):
        for i, (cell, value) in enumerate(zip(tr.cells, values)):
            set_cell(cell, value, bold=first_bold and i == 0)
    keep_rows_whole(table)


# ------------------------------------------------------------------ content

def build(e: dict) -> Document:
    v2, cat, host = e["v2"], e["v2"]["catalog"], e["v2"]["hosting"]
    bands, res, argo, glider = e["depth_bands"], e["residual"], e["argo"], e["glider"]
    coloc, rng, tchp, text = e["colocation"], e["field_range_test"], e["tchp"], e["text_ingest"]
    before, after, hero, casts = v2["amphan_before"], v2["amphan_after"], v2["hero"], v2["amphan_before_casts"]
    shallow, mid = bands["0-300 m"], bands["300-950 m"]
    empty_pct = f"{100 - res['coverage_percent']:.0f}"
    requests_total = len(json.loads(HOSTING.read_text(encoding="utf-8"))["requests"])

    doc = Document(TEMPLATE)
    cp = doc.core_properties
    cp.author = cp.last_modified_by = AUTHOR
    cp.title = "VVater: SIH 2026 idea proposal, PS 26067"

    # header table
    head = doc.tables[0]
    for row, key in zip(head.rows, ["Problem Statement ID", "Problem Statement Title",
                                    "Team ID", "Team Name"]):
        set_cell(row.cells[1], TITLE_FIELDS[key], size=11)

    strip_gray(doc)
    drop_empty_paragraphs(doc)
    for label in ("Outcomes:", "Stretch Goal:", "Research References:"):
        find(doc, label).paragraph_format.keep_with_next = True

    # ---------------------------------------------------------------- 1
    exec_head = find(doc, "1. Executive Summary")
    set_text(exec_head, "1. Executive Summary")
    Writer(exec_head).p(
        "VVater is a browser-native 3D ocean platform for INCOIS. A user picks a box anywhere "
        f"on Earth, a day and one of {cat['variables']} model variables ({cat['from_first_day']} "
        f"of them daily since {cat['first_day'][:4]}, all to the forecast horizon), and the "
        "model's own depth levels are lifted out "
        "as a block that can be cut open and seen from the side. Argo floats stand inside "
        "it, each with its QC flag, data mode and source file, and one click compares a "
        "float with the model on its own day. Data is read where Copernicus and INCOIS publish "
        "it: no archive to build. It is built and measured: against an independent "
        f"glider, the INCOIS analysis has {n(mid['glider']['rmse'])} °C RMSE below 300 m but "
        f"{n(shallow['glider']['rmse'])} °C in the upper 300 m, the layer cyclones feed "
        "on: hidden in a pooled number, plain in 3D. It serves forecasters, fishermen, "
        "students and the public.")

    # ---------------------------------------------------------------- 2.1
    Writer(find(doc, "2.1. Problem Definition")).p(
        "INCOIS generates three-dimensional model fields of temperature, salinity, currents and "
        "biogeochemistry, and receives profiles from Argo floats and gliders. The problem "
        "statement puts the gap plainly: “no integrated, web-based 3D visualization "
        "platform currently exists that can simultaneously render model fields and in-situ "
        "instrument observations in a single interactive environment”, so forecasters "
        "“are therefore forced to toggle between disparate software packages, making it "
        "difficult to rapidly correlate model predictions with observational evidence.”"
    ).p(
        "Why it matters now: the theme is Disaster Management, and the part of the ocean that "
        "decides a Bay of Bengal cyclone's intensity is the warm upper layer, which a map of the "
        "surface cannot show and a single depth slice shows only one level at a time. The brief "
        "lists what this delays: “timely hazard assessment, search-and-rescue support, "
        "fishery advisories, climate monitoring”."
    ).p(
        "What we found building it (the first three are evaluation-harness results, the last "
        "a probe of the links):"
    ).bullets([
        ("The model needs checking too.", f"{rng['failed']} of {rng['checked_cells']:,} cells "
         f"of the INCOIS Bay of Bengal analysis read {n(rng['failed_range'][0])}–"
         f"{n(rng['failed_range'][1])} °C at depth and fail the published Argo global range "
         "test. A renderer that trusts its input paints a heatwave that never happened."),
        ("Observations cover very little of the model.", f"In a ±{e['pairing_days']:.0f}-day "
         f"window only {res['cells_filled']} of {res['cells_total']:,} analysis cells "
         f"({res['coverage_percent']} %) contain any measurement at all."),
        ("A pooled error hides where the model is wrong.", "Against the assimilated Argo "
         f"floats the analysis looks close ({n(coloc['mean_rmse_degC'])} °C RMSE); against an "
         f"independent glider it is {n(coloc['glider_mean_rmse_degC'])} °C, and banding by "
         f"depth puts almost all of that in the upper 300 m ({n(shallow['glider']['rmse'])} °C "
         f"RMSE, bias {n(shallow['glider']['bias'])} °C). One glider deployment "
         f"({glider['deployments'][0].split('-')[0]}), one fortnight off Sri Lanka in 2018: a "
         "finding about that water, not about the analysis everywhere."),
        ("The brief's own data links are dead.", "Both the Argo and glider links are ftp:// "
         "and port 21 is blocked; every source we use is an HTTPS path we found and tested."),
    ]).p(
        "Those most affected are INCOIS's operational forecasters, and through their advisories "
        "coastal communities and fishermen; the brief also names students, the public and "
        "policymakers, for whom model output is “typically inaccessible”.")

    # ---------------------------------------------------------------- 2.2
    nat = find(doc, "National:")
    set_text(nat, "National:", bold=True)
    nat.add_run(" INCOIS and ISRO already publish rich ocean data; what is missing is one "
                "3D scene where model and instruments meet.").font.size = Pt(BODY_PT)
    fill_table(doc.tables[1], [
        ["INCOIS Digital Ocean (MoES, launched Dec 2020)",
         "As announced: a web environment for data integration, “3D and 4D (3D in space "
         "with time animation) data visualization”, analysis and download of in-situ, "
         "satellite and model data on a georeferenced 3D ocean.",
         "Government platform; brings in-situ, remote-sensing and model data together; the "
         "closest national incumbent.",
         "The address given at launch (do.incois.gov.in) did not resolve from our network on "
         "2026-09-26, so we could not assess it first-hand. The announcement does not describe "
         "model–float co-location or QC-flag display, and the 2026 problem statement still "
         "finds that no integrated 3D platform for model fields and in-situ observations "
         "exists."],
        ["INCOIS Live Access Server (las.incois.gov.in)",
         "NOAA PMEL's Live Access Server: on-the-fly plots and custom subsets of gridded model "
         "output, drawn server-side with Ferret; OPeNDAP access.",
         "Mature and reliable; serves INCOIS model output; subsetting and download.",
         "Static 2D plots one request at a time; no interactive 3D volume; drawing instrument "
         "profiles in the same view as the model is not found in its documentation."],
        ["INCOIS ERDDAP (erddap.incois.gov.in)",
         "Data server: 15 gridded and 2 tabular datasets with subsetting, graphs and WMS.",
         "Open, no credentials; the Argo analysis ships a per-cell error field.",
         "A data service, not a viewer: graphs are 2D, one variable at a time, so the error "
         "field is a separate plot, never combined with the field it qualifies; no 3D view."],
        ["MOSDAC (ISRO Space Applications Centre)",
         "Satellite ocean products (surface currents, sea-surface salinity, ocean subsurface, "
         "eddies) with downloads and image galleries.",
         "National satellite products, openly listed.",
         "Product downloads and galleries; an interactive 3D view or an Argo or glider overlay "
         "was not found on the portal."],
    ])
    glob = find(doc, "Global:")
    set_text(glob, "Global:", bold=True)
    glob.add_run(" The leading tools are strong at maps or at profiles, and none puts both "
                 "in one cut-open 3D block.").font.size = Pt(BODY_PT)
    fill_table(doc.tables[2], [
        ["Copernicus MyOcean Pro Viewer (EU)",
         "Web viewer over the whole Copernicus catalogue: depth and time navigation, vertical "
         "profile, transect and time-series graphs, compatible in-situ data, a 3D globe option.",
         "Official, complete catalogue, exports, shareable links.",
         "The ocean is a map at one depth at a time; profiles and transects are separate "
         "graphs, not a volume you can walk around or cut; not built for INCOIS products."],
        ["Ocean Data View / webODV (Alfred Wegener Institute)",
         "Desktop software for station profiles, sections and time series; webODV serves Argo "
         "collections online.",
         "The standard tool for profile analysis; imports Argo, WOD, WOCE and more; free for "
         "non-commercial use.",
         "Desktop-first and expert-oriented; 2D sections; model fields and floats are not in "
         "one 3D scene."],
        ["Argovis (University of Colorado Boulder)",
         "REST API and web app to search, co-locate and visualise Argo, ship profiles, drifters "
         "and gridded products.",
         "Fast, FAIR, a co-location API; we use it as our Argo fallback.",
         "Maps and profile charts in 2D; no 3D volume of an ocean model."],
        ["earth.nullschool.net",
         "Animated globe of surface currents and SST from OSCAR, Copernicus and NOAA products.",
         "Striking and intuitive; widely used for outreach.",
         "Surface only, no depth; no observations; states “no guarantee of accuracy”."],
    ])

    # ---------------------------------------------------------------- 3.1
    Writer(find(doc, "3.1. Solution Overview")).p(
        "VVater lifts a block out of the ocean model, anywhere and on any day, and sets it on "
        "the sea surface so it can be seen from the side, the one direction a map cannot show. "
        "Its faces are sections painted from the model's native depth levels; any side can be "
        "moved inwards to look inside. Argo floats stand in the block as vertical sticks and "
        "click through to a profile against the model. Around it the globe is painted with the "
        "same variable and colour bar, with that day's currents and winds moving over it. An "
        "assistant builds views on request, and a six-lesson course turns the same data into "
        "teaching."
    ).figure(FIGURES / "shot-float.png",
             f"Figure 1. The running build: the Bay of Bengal on {before['day']}, before Cyclone "
             f"Amphan, cut open on the model's {hero['levels']} native levels. One of "
             f"{casts['found']} Argo floats is clicked: the probe shows its data mode, QC count "
             "and source file; the profile compares it with the model.")

    # ---------------------------------------------------------------- 3.2
    for starts in ("Objective 1", "Objective 2", "Objective 3"):
        p = find(doc, starts)
        p._p.getparent().remove(p._p)
    Writer(find(doc, "3.2. Core Objectives")).p(
        "Build a cube generator that serves any box, any day in each variable's coverage and "
        f"any of the catalogue's {cat['variables']} variables ({cat['from_first_day']} of them "
        f"from {cat['first_day'][:4]}) on the model's native depth levels only, and opens a "
        f"prepared cube in under a second. Done: {before['open_seconds_disk_cache']} s for the "
        f"Amphan cube from the disk cache, {before['payload_mb']} MB to the browser.",
        "Objective 1:"
    ).p(
        "Draw every in-situ observation inside the model with its QC flag, data mode and source "
        "file, never dropping a failed reading, and co-locate it with the model level by level. "
        f"Done: {coloc['profiles_compared']} Argo profiles and "
        f"{coloc['glider_casts_compared']} glider casts co-located; {argo['levels_rejected']} "
        f"of {argo['levels']:,} Argo levels that fail QC are drawn as rejected.",
        "Objective 2:"
    ).p(
        "Give the controls the brief names (variable selector, depth slices and cut planes, "
        "day-by-day timeline, colour bar with palette, range and log scale, opacity, vertical "
        "exaggeration, isosurface) and open interfaces (REST, OGC WMS, NetCDF and delimited "
        f"text). Done: the text parser reproduces the NetCDF path exactly, {text['casts_identical']} "
        f"of {text['casts_netcdf']} casts identical.",
        "Objective 3:"
    ).p(
        "By the end of the grand finale, close the known gaps scheduled in section 5.1 and put the "
        "platform in front of INCOIS forecasters, recording what they could and could not do "
        "unassisted. In progress.",
        "Objective 4:")

    # ---------------------------------------------------------------- 3.3
    for starts in ("Novelty 1", "Novelty 2", "Novelty 3"):
        p = find(doc, starts)
        p._p.getparent().remove(p._p)
    Writer(find(doc, "3.3. Novelty and Innovation")).p(
        "Every ocean viewer we examined shows the ocean as a map, one depth at a time. VVater "
        "cuts a block of any size out of the model and looks at it from the side, so a "
        "thermocline, an oxygen-poor layer or the edge of a current is seen, not inferred from a "
        "stack of maps. The block carries the model's native levels and never more: upsampling "
        "would invent structure in the thermocline, the layer that matters most.",
        "Novelty 1: The ocean as a block you can cut."
    ).p(
        "Floats are not dots on a surface but sticks standing in the model at their true depths. "
        "A reading that fails QC is drawn as rejected, never silently dropped, and one click "
        "compares the float with the model on the float's own day.",
        "Novelty 2: Observations stand inside the model."
    ).p(
        "The model is range-tested before it is drawn: impossible cells are masked and counted "
        "on screen. The INCOIS analysis's own error field is drawn, so uncertain water renders "
        "faint. And observed minus modelled is binned into a residual volume where only measured "
        f"cells are filled: {empty_pct} % of the Bay stays empty, because nobody measured it.",
        "Novelty 3: The model gets quality control too."
    ).p(
        "The platform holds no copy of the ocean. It reads only the cloud-store chunks a box "
        f"touches, so every day from {cat['first_day'][:4]} for {cat['from_first_day']} "
        "variables is one request away and a day seen once opens from the cache. Which model a "
        "day comes from is written on screen, and a forecast is always labelled a forecast.",
        "Novelty 4: Any day, any variable, no archive to build."
    ).p(
        "The assistant builds cubes and sets controls on request, but reaches data only "
        "through our own API, and any number in its reply that cannot be traced to a tool "
        "result is flagged on screen. The course's quiz answers are computed from the block on "
        "screen, never from the language model.",
        "Novelty 5: An assistant that acts, and cannot quietly invent a number.")

    # ---------------------------------------------------------------- 3.4
    for starts in ("Performance:", "Accuracy:", "User Adoption:"):
        p = find(doc, starts)
        p._p.getparent().remove(p._p)
    Writer(find(doc, "3.4. Success Metrics")).p(
        "Every figure here is produced by the repository's evaluation script "
        "(docs/13-eval-results.md); nothing is estimated.", italic=True
    ).p(
        f"Amphan cube ({before['cells'][0]}×{before['cells'][1]} cells × {before['levels']} "
        f"levels) opens in {before['open_seconds_disk_cache']} s from the disk cache, "
        f"{before['payload_mb']} MB to the browser. Over a cold, full session touching every "
        f"feature, {host['requests_ok']} of {requests_total} requests answered, peak memory "
        f"{host['peak_mb']} MB. A cold open waits on Copernicus and is not claimed.",
        "Performance:"
    ).p(
        f"Co-location against the INCOIS analysis: {n(coloc['mean_rmse_degC'])} °C RMSE over "
        f"{coloc['profiles_compared']} Argo profiles (a consistency check: the analysis "
        f"assimilates them) and {n(coloc['glider_mean_rmse_degC'])} °C over "
        f"{coloc['glider_casts_compared']} independent glider casts. Target, met: independent "
        "error above the fit residual, as theory predicts; an inverted depth axis or swapped "
        "lat/lon would give errors of order 10 °C.",
        "Accuracy:"
    ).p(
        f"{argo['levels_rejected']} of {argo['levels']:,} Argo levels fail QC and all are shown "
        f"as rejected; {rng['failed']} impossible model cells are masked and counted; the text "
        f"and NetCDF parsers give identical casts, {text['casts_identical']} of "
        f"{text['casts_netcdf']} ({text['levels_identical']:,} levels).",
        "Data integrity:"
    ).p(
        "Not yet measured. Target for the finale: INCOIS forecasters complete three tasks "
        "unassisted (build a cube, open a float against the model, find where the model is "
        "least certain), with completion and time recorded.",
        "User adoption:")

    # ---------------------------------------------------------------- 4.1
    Writer(find(doc, "4.1. Technical Approach")).p(
        "Two tiers. A static TypeScript viewer on CesiumJS runs in any modern browser with "
        "nothing installed. One Python process (FastAPI) serves cubes, floats, surface layers, "
        "wind, fishing advisories and the assistant as JSON and raw float32 arrays, plus OGC WMS "
        "for the Bay layers. Data is fetched on demand from its publishers and cached by chunk "
        "on disk."
    ).p("From request to picture:", bold=True).bullets([
        ("Ask.", "A box, a day and a variable, drawn by hand, picked from nine named scenarios "
         "or asked of the assistant. The URL carries the whole view."),
        ("Pick the model for that day.", "The catalogue resolves each day to the reanalysis "
         "where it exists and the analysis & forecast after it; forecast days are labelled."),
        ("Read only what the box touches.", "The Copernicus ARCO zarr chunks inside the box are "
         "read from S3 and cached on disk; nothing else is downloaded."),
        ("Check and shape.", "Values outside the variable's published range are masked and "
         "counted; density and sound speed are derived with TEOS-10; the block keeps its native "
         "levels and is averaged horizontally to at most 160 cells a side for display, with the "
         "factor stated on screen."),
        ("Paint.", "The browser paints six faces as sections through the data, a sea floor from "
         "bathymetry, and any cut the user makes."),
        ("Floats.", "Core and BGC Argo in the box within ±2 days come from Ifremer's ERDDAP with "
         "every level's QC flag, data mode and file name; temperature is converted to potential "
         "temperature before comparison. Clicking a float co-locates it with the model on its own "
         "day through a tested module."),
    ]).figure(FIGURES / "doc-flow.png",
              "Figure 2. How a cube is made, and what happens to data that fails a test: "
              "masked and counted, or drawn as rejected, never dropped."
    ).p(
        "The INCOIS Bay of Bengal analysis is also drawn as a true voxel volume with its "
        "uncertainty, observation density and residual, over the one date where INCOIS, Argo, a "
        "glider and Copernicus are all present. That is where the depth-banded result below "
        "comes from."
    ).figure(FIGURES / "doc-bands.png",
             "Figure 3. RMSE of the INCOIS analysis by depth band: Argo (grey, assimilated) and "
             "the independent glider (blue). Below 300 m they agree; the error is in the upper "
             "300 m. One glider deployment, so a finding about that water.", width_cm=13)

    # ---------------------------------------------------------------- 4.2
    tech = find(doc, "4.2. Tech Stack")
    tech.paragraph_format.keep_with_next = True
    Writer(tech).figure(FIGURES / "doc-arch.png",
                        "Figure 4. The build in three tiers, with pinned versions.", width_cm=11)
    fill_table(doc.tables[3], [
        ["Frontend", "CesiumJS (@cesium/engine 26.3.0, @cesium/widgets 16.2.0), TypeScript "
         "5.9, Vite 7, no UI framework; Canvas 2D particle overlay for currents and winds; "
         "optional Cesium ion terrain and imagery"],
        ["Backend", "Python 3.14, FastAPI + Uvicorn; xarray, netCDF4, zarr and dask for NetCDF "
         "and cloud stores; numpy, scipy and gsw (TEOS-10); copernicusmarine toolbox; truststore "
         "so TLS is always verified"],
        ["Database", "None. Data stays with its publishers; a disk chunk cache holds what has "
         "been read, and uploaded CSV casts live in memory"],
        ["AI/ML Model(s)", "No trained model. The assistant uses hosted language models through "
         "AIRouter (openai/gpt-4.1-mini, then google/gemini-2.5-flash, then "
         "deepseek/deepseek-v4-flash); it calls our API's tools and nothing else, and untraceable "
         "numbers are flagged"],
        ["Cloud Services", "Copernicus Marine ARCO stores on S3; INCOIS, Ifremer and IOOS "
         "ERDDAP; UCAR THREDDS (GFS wind); NASA GIBS (land relief); static hosting (Vercel) and "
         "one API server"],
        ["Standards", "CF Conventions (read defensively, every assumption recorded), OGC WMS "
         "1.3.0, REST/JSON, Argo QC flags; IOOS QARTOD flags read where present (else "
         "carried as unevaluated)"],
    ])

    # ---------------------------------------------------------------- 4.3
    Writer(find(doc, "4.3. Hardware & Data Requirements")).p(
        "No hardware is built or bought. The client is any modern browser; the server "
        "requirement below is measured, not estimated. Prices are not given because the server "
        "is meant to run on existing INCOIS infrastructure."
    )
    fill_table(doc.tables[4], [
        ["Any computer or tablet with a WebGL 2 browser", "per user", "Nothing to buy"],
        [f"API server or VM: 1 GB RAM minimum (measured peak {host['peak_mb']} MB), 1–2 vCPU, "
         "5–10 GB persistent disk for the chunk cache", "1", "Existing infrastructure"],
        ["Projector or large display for outreach (optional)", "per venue", "Existing equipment"],
    ], first_bold=False, footer=1)
    total = doc.tables[4].rows[-1]
    set_cell(total.cells[0], "Total cost: no new hardware. One server-class VM on existing "
             "infrastructure.", bold=True)
    new_para_after(doc.tables[4], "Datasets, all read over HTTPS from their "
                               "publishers, licences as each publisher states them:", bold=True)
    fill_table(doc.tables[5], [
        ["Copernicus Marine GLORYS12 reanalysis and global analysis & forecast, physics and "
         "biogeochemistry (ARCO zarr)",
         "data.marine.copernicus.eu/product/GLOBAL_MULTIYEAR_PHY_001_030",
         "Copernicus Marine Service licence: free and open, attribution required"],
        ["INCOIS Argo 10-day Variational Analysis (incois_argo_10d_VAM), with per-cell error",
         "erddap.incois.gov.in/erddap", "Free use and redistribution, no warranty (dataset "
         "licence attribute)"],
        ["Argo core and BGC floats", "erddap.ifremer.fr/erddap; data-argo.ifremer.fr",
         "Argo open data policy: freely available; acknowledgement and the Argo DOI requested"],
        ["U.S. IOOS Glider DAC, deployment ru29-20180812T0220", "gliders.ioos.us/erddap",
         "“May be redistributed and used without restriction” (dataset licence)"],
        ["INCOIS Potential Fishing Zone advisories", "incois.gov.in/MarineFisheries/"
         "TextDataHome?mfid=1", "Public advisories, shown as published and attributed"],
    ], first_bold=False)
    Writer(doc.tables[5]).p(
        "Collection strategy: nothing is copied in advance. Each request reads the chunks or "
        "rows it needs and caches them on disk; every observation keeps its QC flag, data mode "
        "and source file. The brief's two FTP links are dead from a normal network (port 21), "
        "so the Argo GDAC HTTPS mirror and the IOOS Glider DAC replace them. Wind (NCEP GFS via "
        "UCAR) and land relief (NASA GIBS) are display layers, attributed on screen.")

    # ---------------------------------------------------------------- 5.1
    fill_table(doc.tables[6], [
        ["Before the finale: built and measured",
         f"Cube generator over {cat['variables']} variables; floats inside the model with QC; co-location and "
         "the evaluation harness; four views, immersive and cinematic; winds, waves and fishing "
         "advisories; the assistant; the six-lesson course; INCOIS Bay volume with uncertainty "
         "and residual; OGC WMS. Every number in this proposal comes from the harness.",
         "Complete, whole team"],
        ["Hours 0–6: floats readable at a glance",
         "Colour each float stick by its values even behind a face (today it draws as white "
         "dashes there); the assistant reports the depth it actually built; drawing a box "
         "clears the old cube. ACCEPTED: a misfit is visible without opening a chart; no "
         "reply states a depth the viewer did not build.",
         "6 h, Frontend + Assistant"],
        ["Hours 6–12: the float window, measured",
         "The harness bins float-minus-model misfit by time separation and picks the window "
         "where it stops growing, replacing today's provisional ±2 days. ACCEPTED: a "
         "cube-wide misfit number in docs/13-eval-results.md.",
         "6 h, Lead + Ocean science"],
        ["Hours 12–18: the brief's own glider archive",
         "Index the EGO/Coriolis glider catalogue (the archive the dead FTP link served) by "
         "position, and read its gliders with their QC. ACCEPTED: a second glider source in "
         "the viewer and in the harness.",
         "6 h, Ingestion"],
        ["Hours 18–24: cyclone heat potential as a map",
         "Integrate heat above 26 °C over every grid column of the analysis, as a globe layer "
         "and a WMS layer. ACCEPTED: the map matches the per-float values the harness already "
         "checks.",
         "6 h, Ocean science + Ingestion"],
        ["Hours 24–30: host and frame rate",
         "Move the API to a host with 1 GB or more; measure frame rates on the demo machine "
         "and the venue display and record them. ACCEPTED: the live site answers with the "
         "laptop off; measured fps written down.",
         "6 h, QA and deploy"],
        ["Hours 30–36: INCOIS walk-through and pitch",
         "Forecasters or mentors run the three tasks in 3.4 unassisted; record completion, "
         "time and every hesitation, and fix the top friction point. ACCEPTED: results "
         "reported as measured, whatever they are.",
         "6 h, whole team"],
    ])

    # ---------------------------------------------------------------- 5.2
    fill_table(doc.tables[7], [
        ["Mr. A Mohamed Armaan", "Team Lead: data pipeline and evaluation",
         "The cube generator and its native-levels rule; the co-location module; the "
         "evaluation harness every number comes from. Planning and review."],
        ["Ms. Pooja Shree Ravichandar", "Frontend: 3D viewer",
         "The CesiumJS scene: the cube's painted faces and cuts, the four views and cameras, "
         "colour bar and layer controls, the float sticks and profile panel."],
        ["Ms. Aashisha Wincy A", "Ocean science and outreach",
         "Argo QC conventions, TEOS-10 conversions and cyclone heat potential; the six-lesson "
         "course and its sources; liaison with INCOIS scientists."],
        ["Mr. Jerem Agsae Jebaz", "Assistant and language",
         "The tool-calling assistant, its whitelisted viewer actions and the number check; the "
         "user guide it reads; the course in Indian languages."],
        ["Mr. Denny Mathew", "Data ingestion and interoperability",
         "NetCDF and delimited-text parsers, CF normalisation, the ERDDAP, ARCO and glider "
         "sources, the fishing advisory reader and the OGC WMS endpoint."],
        ["Mr. Jonathan Jerald Eapen", "QA, hosting and deploy",
         "Tests and fixtures, the hosting measurement, deploys, graphics tiers and frame-rate "
         "measurement."],
    ])

    # ---------------------------------------------------------------- outcomes
    outcomes = find(doc, "Outcomes:")
    Writer(outcomes).p(
        "VVater runs end to end today, and its source, documents and measured numbers are "
        f"public at {REPO}. The live site currently depends on a development machine for its "
        "API; section 5.1 moves it to a proper host. The outcome it is built for is faster, "
        "better-founded advisories: a forecaster sees the warm layer a cyclone feeds on in 3D "
        "and sees where the model is uncertain or unmeasured before trusting it. Between the "
        f"two Amphan cubes the Bay's mean surface temperature falls by {v2['amphan_cooling_c']} °C, "
        "and the cold wake is visible on one colour bar."
    ).bullets([
        ("Forecasters:", "model and instruments in one scene instead of several programs."),
        ("Fishermen:", "INCOIS's own fishing-zone advisories for all fourteen sectors, beside "
         "that day's waves and wind; our own indicator is labelled not an advisory."),
        ("Students and the public:", "a six-lesson course for classes 8–12 on real data, and an "
         "immersive view for exhibitions."),
        ("Economics:", "open source, no licences, no archive to buy or store; one server with "
         f"1 GB of memory measured sufficient ({host['peak_mb']} MB peak)."),
    ]).p(
        "We have not measured time saved per advisory and will not estimate it; the finale "
        "walk-through is the first step towards that number.")

    Writer(find(doc, "Stretch Goal:")).p(
        "Stream the volume in tiles so full 1/12° resolution and long time animations fit in "
        "the browser; particle trajectories integrated through a sequence of days rather than "
        "one day's flow; a depth-aware model range test agreed with INCOIS; the course in "
        "Hindi, Tamil and Bengali; more sensors (CTD, moorings, HF radar, ADCP) through the "
        "same parser table the Argo, glider and CSV paths already share.")

    refs = [
        "SIH 2026 Problem Statement 26067, INCOIS / MoES, quoted verbatim — "
        f"{REPO}/blob/main/docs/01-problem-statement.md",
        "Press Information Bureau (2020-12-30). Launch of INCOIS ‘Digital Ocean’ — "
        "pib.gov.in/PressReleseDetailm.aspx?PRID=1684418",
        "NOAA PMEL, Live Access Server — ferret.pmel.noaa.gov/LAS",
        "INCOIS ERDDAP — erddap.incois.gov.in/erddap",
        "MOSDAC, Space Applications Centre, ISRO — mosdac.gov.in",
        "Copernicus Marine, MyOcean Pro features — help.marine.copernicus.eu/en/articles/"
        "4794675-overview-of-myocean-pro-features",
        "Schlitzer R., Ocean Data View — odv.awi.de",
        "Argovis: Building a FAIR Ocean Data Service. J. Atmos. Oceanic Technol. 42(11), 2025 — "
        "argovis.colorado.edu",
        "earth.nullschool.net, About — earth.nullschool.net/about.html",
        "Copernicus Marine GLORYS12 (GLOBAL_MULTIYEAR_PHY_001_030) — data.marine.copernicus.eu",
        "Argo GDAC and Ifremer ERDDAP — data-argo.ifremer.fr; erddap.ifremer.fr/erddap",
        "U.S. IOOS Glider DAC — gliders.ioos.us/erddap",
        "INCOIS Potential Fishing Zone advisories — incois.gov.in/MarineFisheries/"
        "TextDataHome?mfid=1",
        "Wong A. et al., Argo Quality Control Manual — doi.org/10.13155/33951",
        "U.S. IOOS QARTOD manuals — ioos.noaa.gov/project/qartod",
        "CF Conventions for NetCDF — cfconventions.org",
        "OGC Web Map Service 1.3.0 — ogc.org/standards/wms",
        "IOC, SCOR and IAPSO (2010), TEOS-10 — teos-10.org",
        "Leipper D. & Volgenau D. (1972), Hurricane heat potential of the Gulf of Mexico, "
        "J. Phys. Oceanogr. 2 — doi.org/10.1175/1520-0485(1972)002<0218:HHPOTG>2.0.CO;2",
        "Thyng K. et al. (2016), True colors of oceanography, Oceanography 29(3) — "
        "doi.org/10.5670/oceanog.2016.66",
        "CesiumJS — cesium.com/platform/cesiumjs",
        "India Meteorological Department, report on Super Cyclonic Storm Amphan (2020). "
        "The Amphan dates.",
        "Repository: source, documents and every measured number — " + REPO,
    ]
    w = Writer(find(doc, "Research References:"))
    for i, ref in enumerate(refs, 1):
        w.p(ref, f"{i}.", size=10)
    return doc


# ------------------------------------------------------------------ check

LEFTOVERS = ["SIHxxxx", "K26xxx", "[Name", "e.g.,", "To develop", "To implement",
             "To integrate", "Instructions:", "[Explain", "Day 1: Setup"]


def words_after(doc, starts: str) -> int:
    ps = doc.paragraphs
    i = next(k for k, p in enumerate(ps) if p.text.strip().startswith(starts))
    return len(ps[i + 1].text.split())


def check() -> None:
    doc = Document(OUT)
    everything = "\n".join(p.text for p in doc.paragraphs)
    for t in doc.tables:
        everything += "\n".join(c.text for row in t.rows for c in row.cells)
    problems = [f"leftover template text {s!r}" for s in LEFTOVERS if s in everything]
    gray = [r.text for p in doc.paragraphs for r in p.runs
            if r.font.color is not None and r.font.color.type and str(r.font.color.rgb) == GRAY]
    problems += [f"gray run left: {g[:40]!r}" for g in gray]
    summary, overview = words_after(doc, "1. Executive Summary"), words_after(doc, "3.1.")
    if summary >= 150:
        problems.append(f"executive summary is {summary} words; the template allows under 150")
    if overview > 120:
        problems.append(f"solution overview is {overview} words; the template asks for about 100")
    for p in problems:
        print("FAIL", p)
    if problems:
        raise SystemExit(1)
    print(f"ok: executive summary {summary} words, solution overview {overview} words, "
          f"{len(doc.inline_shapes)} figures")


def main() -> None:
    if "--check" in sys.argv:
        return check()
    if not TEMPLATE.exists():
        raise SystemExit(f"missing {TEMPLATE.name} in the repo root (the portal's file)")
    e = json.loads(EVAL.read_text(encoding="utf-8"))
    diagrams(e)
    OUT.parent.mkdir(parents=True, exist_ok=True)
    build(e).save(OUT)
    print(f"wrote {OUT.relative_to(ROOT)}")
    check()


if __name__ == "__main__":
    main()
