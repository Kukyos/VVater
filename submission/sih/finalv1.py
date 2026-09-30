"""finalv1: the submitted deck brought level with the build (docs/24-brief-coverage.md)."""
import copy, sys
from pptx import Presentation
from pptx.util import Inches

SRC = "submission/sih/final/VVater-SIH2026-final.pptx"   # as submitted, 27 Sept
OUT = "submission/sih/final/VVater-SIH2026-finalv1.pptx"
p = Presentation(SRC)
S2, S3, S4, S6 = (p.slides[i] for i in (1, 2, 3, 5))

def get(s, i):
    return next(sh for sh in s.shapes if sh.shape_id == i)

def runs(s, i, para=0):
    return get(s, i).text_frame.paragraphs[para].runs

def set_runs(s, i, texts, para=0):
    rs = runs(s, i, para)
    assert len(rs) == len(texts), (i, [r.text for r in rs])
    for r, t in zip(rs, texts):
        r.text = t

def by_text(s, text):
    return next(sh for sh in s.shapes if sh.has_text_frame and sh.text_frame.text.startswith(text))

next_id = [10_000]
def clone(s, i, dy):
    """Copy shape i on slide s, dy inches lower, with a fresh id."""
    el = copy.deepcopy(get(s, i)._element)
    s.shapes._spTree.append(el)
    for c in el.iter():
        if c.tag.endswith("}cNvPr"):
            next_id[0] += 1
            c.set("id", str(next_id[0]))
    new = s.shapes[-1]
    new.top = new.top + Inches(dy)
    return new

# ---- slide 2
set_runs(S2, 82, ["Instruments:", " Argo and BGC floats, gliders, RAMA moorings, CTD casts, each with its QC state and source file"])
set_runs(S2, 91, ["23 variables, cut planes, time-step animation (Play), palette, range and log scale"])
set_runs(S2, 93, ["a variable or ML product is one catalogue line; a sensor, one registry entry"])
set_runs(S2, 97, ["Isosurface (the 20 °C isotherm, Bay volume) · layer opacity · vertical exaggeration · NetCDF and text parsers · REST API · OGC WMS (Bay layers) · ADCP and HF-radar readers"])
set_runs(S2, 100, ["click one: its readings beside the model's, on the float's own day"])

# ---- slide 3: the assistant's real models, and the in-situ TAC as a data source
tb = get(S3, 133).text_frame.paragraphs
for para, t in zip(tb, ["AI models, via AIRouter", "GPT-4.1 mini; Gemini, DeepSeek", "as fallbacks; questions capped"]):
    assert len(para.runs) == 1
    para.runs[0].text = t
cop = by_text(S3, "Copernicus Marine")
cop.text_frame.paragraphs[1].runs[0].text = "ocean model and ML products, 1993 on"
assert len(cop.text_frame.paragraphs[1].runs) == 1
clone(S3, 59, 0.72)                      # the frame under NOAA
tac = clone(S3, 130, 2.16)               # Ifremer's text, three boxes down
tac.text_frame.paragraphs[0].runs[0].text = "Copernicus In Situ TAC"
tac.text_frame.paragraphs[1].runs[0].text = "moorings, CTD, ADCP, HF-radar"

# ---- slide 4
set_runs(S4, 112, ["Free public data; only AI calls cost"], para=1)
set_runs(S4, 114, ["built, live, ML"], para=1)
set_runs(S4, 117, ["HF-radar, more basins"], para=1)

# ---- slide 6: one data source and one method added, then renumbered
nums = {sh.shape_id: int(sh.text_frame.text) for sh in S6.shapes
        if sh.has_text_frame and sh.text_frame.text.strip().isdigit()}
DY_DATA = 9.15 - 3.18                    # item 1's rows, moved below item 8
n9 = clone(S6, 24, DY_DATA)
t9 = clone(S6, 25, DY_DATA)
u9 = clone(S6, 26, DY_DATA)
clone(S6, 27, DY_DATA + 0.13)          # below the URL's second line
t9.text_frame.paragraphs[0].runs[0].text = "Copernicus Marine In Situ TAC and MULTIOBS 3D BGC. "
t9.text_frame.paragraphs[0].runs[1].text = "Moorings, CTD, ADCP, HF-radar; the ML products."
u = u9.text_frame.paragraphs[0].runs[0]
u.text = "data.marine.copernicus.eu/product/INSITU_GLO_PHYBGCWAV_DISCRETE_MYNRT_013_030"
u.hyperlink.address = "https://data.marine.copernicus.eu/product/INSITU_GLO_PHYBGCWAV_DISCRETE_MYNRT_013_030"
DY_STD = 8.66 - 7.98                     # CesiumJS's rows, one pitch lower
n18 = clone(S6, 102, DY_STD)
t18 = clone(S6, 103, DY_STD)
u18 = clone(S6, 104, DY_STD)
clone(S6, 100, DY_STD + 0.69)            # the rule under the new item
t18.text_frame.paragraphs[0].runs[0].text = "Sauzède et al. (2016), J. Geophys. Res. Oceans. "
t18.text_frame.paragraphs[0].runs[1].text = "The ML method."
u = u18.text_frame.paragraphs[0].runs[0]
u.text = "doi.org/10.1002/2015JC011408"
u.hyperlink.address = "https://doi.org/10.1002/2015JC011408"
for i, n in nums.items():                # 9-16 -> 10-17, 17-23 -> 19-25
    get(S6, i).text_frame.paragraphs[0].runs[0].text = str(n + (1 if n <= 16 and n >= 9 else 2 if n >= 17 else 0))
    if n == 9:
        get(S6, i).width = get(S6, 72).width
n9.text_frame.paragraphs[0].runs[0].text = "9"
n18.text_frame.paragraphs[0].runs[0].text = "18"

p.save(sys.argv[1] if len(sys.argv) > 1 else OUT)
print("wrote", OUT)
