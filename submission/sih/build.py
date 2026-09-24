"""Build the SIH idea-submission deck from the official 2026 template.

    python submission/sih/figures.py    # eval-latest.json -> figures/body-*.html
    python submission/sih/render.py     # figures/body-*.html -> figures/body-*.png
    python submission/sih/build.py      # -> submission/sih/final/VVater-SIH2026.pptx

Reads the template the portal supplies; never modifies it.

The template's own masthead, title, team badge, SIH mark, footer bar and page number are
kept, because the portal requires the provided template. Everything between the masthead
and the footer is replaced with a single full-bleed composition per slide, designed in
HTML and rendered as one image by render.py. That is the whole trick: the template
supplies the chrome, and the slide body gets typeset properly instead of being poured
into a placeholder textbox.

Consequence worth knowing: body text on slides 2-6 is part of an image, so it is not
selectable in the PDF. It is also not at the mercy of whichever fonts are installed on
the machine that opens the file, which for a submission is the trade worth making.
Rewording is an edit to figures.py and a rerun.
"""

import sys
from pathlib import Path

from pptx import Presentation
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parents[2]
HERE = Path(__file__).resolve().parent
FIGURES = HERE / "figures"
TEMPLATE = ROOT / "SIH2026-IDEA-Presentation-Format.pptx"
# An optional path argument writes elsewhere, e.g. while the deck is open in PowerPoint.
OUT = Path(sys.argv[1]).resolve() if len(sys.argv) > 1 else HERE / "final" / "VVater-SIH2026.pptx"

# The body area, in inches on the 13.33 x 7.5 slide: below the masthead, above the
# footer bar, bleeding to both edges. figures/deck.css renders to exactly this shape.
BODY_TOP = 1.25
BODY_HEIGHT = 5.70
SLIDE_WIDTH = 13.333

# Slide index -> the figure that replaces its body. Slide 0 is the title page and keeps
# its own text; slide 6 is the portal's "important instructions" page and is dropped,
# because the portal allows six slides in total.
BODIES = {
    1: "body-solution",
    2: "body-technical",
    3: "body-feasibility",
    4: "body-impact",
    5: "body-references",
}
DROP = [6]

# The team is Team Null; VVater is the project, carried by the logo and the idea title.
TEAM_NAME = "Team Null"
LOGO = FIGURES / "logo.png"   # the viewer's favicon, rendered by render.py

# The title page, laid out like the team's PS 26047 deck: the same six lines, bulleted,
# Arial Bold in black, "Label - value". Values verbatim from docs/01-problem-statement.md.
TITLE_FIELDS = {
    "Problem Statement ID": "SIH26067",
    "Problem Statement Title": (
        "Develop a web-based interactive 3D visualization platform that integrates "
        "numerical ocean model outputs and in-situ observations."
    ),
    "Theme": "Disaster Management",
    "PS Category": "Software",
    "Team ID": "",
    "Team Name": TEAM_NAME,
}


def strip_body(slide) -> None:
    """Remove the template's instruction textbox, keeping every other shape.

    Identified by being a plain text box rather than a placeholder or picture: the
    masthead title, the page number and the footer are placeholders, the SIH mark and
    the team oval are a picture and an auto shape, so the only text box on a content
    slide is the instruction block we are replacing.
    """
    for shape in list(slide.shapes):
        if shape.shape_type == 17 and shape.has_text_frame:  # 17 = TEXT_BOX
            shape._element.getparent().remove(shape._element)


def place_body(slide, image: Path) -> None:
    """Drop one full-bleed image into the body area and push it behind nothing."""
    slide.shapes.add_picture(
        str(image),
        Emu(0), Inches(BODY_TOP),
        width=Inches(SLIDE_WIDTH), height=Inches(BODY_HEIGHT),
    )


# Slide 2's masthead placeholder reads "IDEA TITLE" in the template; it is ours to fill.
# Short on purpose: the masthead holds about twenty characters at the template's size
# (a full sentence ran off the top of the slide). The sentence is the slide's lead line.
IDEA_TITLE = "VVater: Ocean in 3D"


def fill_idea_title(slide) -> None:
    for shape in slide.shapes:
        if shape.is_placeholder and shape.has_text_frame and "IDEA TITLE" in shape.text_frame.text:
            runs = shape.text_frame.paragraphs[-1].runs or [shape.text_frame.paragraphs[0].add_run()]
            for para in shape.text_frame.paragraphs:
                for run in para.runs:
                    run.text = ""
            runs[0].text = IDEA_TITLE
            return
    raise SystemExit("no IDEA TITLE placeholder on slide 2; the template changed")


def place_logo(slide) -> None:
    """Every content slide carries a "Your Team Name" oval top left; it becomes the logo,
    square, centred where the oval was."""
    for shape in list(slide.shapes):
        if shape.has_text_frame and "Your Team Name" in shape.text_frame.text:
            size = shape.height
            slide.shapes.add_picture(str(LOGO), shape.left + (shape.width - size) // 2,
                                     shape.top, width=size, height=size)
            shape._element.getparent().remove(shape._element)


def fill_title(slide) -> None:
    """The six title fields, styled as on the PS 26047 deck (scaled to this slide)."""
    from pptx.dml.color import RGBColor
    from pptx.oxml.ns import qn

    for shape in slide.shapes:
        if not (shape.has_text_frame and "Problem Statement ID" in shape.text_frame.text):
            continue
        frame = shape.text_frame
        frame.clear()
        frame.word_wrap = True
        for index, (label, value) in enumerate(TITLE_FIELDS.items()):
            para = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
            para.space_after = Pt(14)
            pPr = para._p.get_or_add_pPr()
            pPr.set("marL", str(Inches(0.3)))
            pPr.set("indent", str(-Inches(0.3)))
            pPr.append(pPr.makeelement(qn("a:buFont"), {"typeface": "Arial"}))
            pPr.append(pPr.makeelement(qn("a:buChar"), {"char": "•"}))
            run = para.add_run()
            run.text = f"{label} - {value}" if value else f"{label} - "
            run.font.size = Pt(20)
            run.font.bold = True
            run.font.name = "Arial"
            run.font.color.rgb = RGBColor(0, 0, 0)
        return
    raise SystemExit("no title field box on slide 1; the template changed")


def main() -> None:
    if not TEMPLATE.exists():
        raise SystemExit(
            f"missing {TEMPLATE.name}. The portal template is not committed (it is not "
            "ours to redistribute) -- drop it in the repo root. See docs/16-submission.md."
        )

    missing = [n for n in BODIES.values() if not (FIGURES / f"{n}.png").exists()]
    if missing:
        raise SystemExit(
            f"missing rendered figures: {missing}. Run figures.py then render.py first."
        )

    deck = Presentation(str(TEMPLATE))

    fill_title(deck.slides[0])
    fill_idea_title(deck.slides[1])
    for slide in deck.slides:
        place_logo(slide)

    for index, name in BODIES.items():
        slide = deck.slides[index]
        strip_body(slide)
        place_body(slide, FIGURES / f"{name}.png")

    # Removing by index from the end, so earlier indices stay valid.
    xml_slides = deck.slides._sldIdLst
    slides = list(xml_slides)
    for index in sorted(DROP, reverse=True):
        xml_slides.remove(slides[index])

    OUT.parent.mkdir(parents=True, exist_ok=True)
    deck.save(str(OUT))
    print(f"wrote {OUT}  ({len(deck.slides)} slides)")


if __name__ == "__main__":
    main()
