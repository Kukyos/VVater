"""Build the deck with every block editable, for import into Canva.

    python submission/sih/figures.py        # figures/body-*.html
    python submission/sih/build_canva.py    # -> final/VVater-SIH2026-canva.pptx

Then in Canva: Create a design -> Import file -> pick the .pptx.

build.py flattens each slide body into one image. This one rebuilds the same bodies as
native shapes: every text block is a text box, every card and rule a rectangle, the
flowchart boxes, diamonds and arrows are shapes and lines, the bar chart is rectangles,
and only the screenshots stay pictures.

It does not re-typeset anything by hand. Each body page is laid out by headless Chrome
with Arial swapped in (the sans present both here and in Canva, so the widths Chrome
measures are the widths PowerPoint and Canva get), a script reads back the position and
style of every box, text run, image and svg element, and this file places a shape for
each at the same spot. Rewording is still an edit to figures.py and a rerun.
"""

import json
import re
import subprocess
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

from build import BODIES, BODY_TOP, DROP, TEMPLATE, fill_idea_title, fill_title, place_logo, strip_body
from render import CHROME

HERE = Path(__file__).resolve().parent
FIGURES = HERE / "figures"
OUT = HERE / "final" / "VVater-SIH2026-canva.pptx"

DPI = 300              # a body page is 4000 x 1710 px = 13.33 x 5.70 in
PT = 72 / DPI          # px -> pt
FONT, MONO = "Arial", "Courier New"

# Injected into a copy of each body page: Arial everywhere, then a walk of the laid-out
# page that writes what it finds as JSON into the DOM, for --dump-dom to carry out.
PROBE = r"""
<style>* { font-family: Arial, sans-serif !important; }
.mono, .mono *, .url { font-family: "Courier New", monospace !important; }</style>
<script>
window.addEventListener("load", () => {
  const out = [];
  const rgb = c => { const m = /rgba?\(([^)]+)\)/.exec(c || ""); if (!m) return null;
    const p = m[1].split(",").map(Number); return p.length > 3 && p[3] === 0 ? null : p.slice(0, 3); };
  const box = r => [r.left, r.top, r.width, r.height];
  const style = el => { const s = getComputedStyle(el); return { size: parseFloat(s.fontSize),
    bold: parseInt(s.fontWeight) >= 600, italic: s.fontStyle === "italic", color: rgb(s.color),
    mono: /Courier/.test(s.fontFamily), upper: s.textTransform === "uppercase" }; };
  const inline = el => getComputedStyle(el).display === "inline";

  function collect(el, runs, rects) {
    for (const n of el.childNodes) {
      if (n.nodeType === 3) {
        const text = n.textContent.replace(/\s+/g, " ");
        if (!text) continue;
        const range = document.createRange(); range.selectNodeContents(n);
        for (const r of range.getClientRects()) if (r.width > 0) rects.push(r);
        runs.push({ text, ...style(n.parentElement) });
      } else if (n.nodeType === 1 && n.tagName === "BR") runs.push({ text: "\n" });
      else if (n.nodeType === 1 && inline(n)) collect(n, runs, rects);
    }
  }

  function svg(root) {
    const M = root.getScreenCTM();
    const at = (x, y) => { const p = root.createSVGPoint(); p.x = x; p.y = y; const q = p.matrixTransform(M); return [q.x, q.y]; };
    for (const el of root.querySelectorAll("rect, polygon, polyline, line, text")) {
      const s = getComputedStyle(el), sw = parseFloat(s.strokeWidth || 0) * M.a;
      const paint = { fill: rgb(s.fill), stroke: s.stroke === "none" ? null : rgb(s.stroke), sw };
      if (el.tagName === "rect") out.push({ t: "rect", r: box(el.getBoundingClientRect()), ...paint });
      else if (el.tagName === "polygon") out.push({ t: "diamond", r: box(el.getBoundingClientRect()), ...paint });
      else if (el.tagName === "text") {
        const anchor = s.textAnchor;
        out.push({ t: "text", svg: true, r: box(el.getBoundingClientRect()),
          align: anchor === "middle" ? "center" : anchor === "end" ? "right" : "left",
          runs: [{ text: el.textContent, size: parseFloat(el.getAttribute("font-size") || s.fontSize) * M.a,
                   bold: parseInt(s.fontWeight) >= 600, color: rgb(s.fill) }] });
      } else {
        const pts = el.tagName === "line"
          ? [[+el.getAttribute("x1"), +el.getAttribute("y1")], [+el.getAttribute("x2"), +el.getAttribute("y2")]]
          : el.getAttribute("points").trim().split(/\s+/).map(p => p.split(",").map(Number));
        out.push({ t: "line", pts: pts.map(([x, y]) => at(x, y)), color: paint.stroke, sw,
                   arrow: el.hasAttribute("marker-end") });
      }
    }
  }

  function walk(el) {
    const s = getComputedStyle(el);
    if (s.display === "none" || s.visibility === "hidden") return;
    if (el.tagName.toLowerCase() === "svg") return svg(el);
    const r = el.getBoundingClientRect();
    if (!["HTML", "BODY"].includes(el.tagName) && !el.classList.contains("slide")) {
      const bg = rgb(s.backgroundColor);
      if (bg && bg.join() !== "255,255,255") out.push({ t: "rect", r: box(r), fill: bg });
      if (el.tagName === "IMG") out.push({ t: "img", src: el.getAttribute("src"), r: box(r),
        fit: s.objectFit, pos: s.objectPosition, nw: el.naturalWidth, nh: el.naturalHeight });
      for (const side of ["Top", "Right", "Bottom", "Left"]) {
        const w = parseFloat(s["border" + side + "Width"]), c = rgb(s["border" + side + "Color"]);
        if (!(w > 0 && c)) continue;
        const rr = { Top: [r.left, r.top, r.width, w], Bottom: [r.left, r.bottom - w, r.width, w],
                     Left: [r.left, r.top, w, r.height], Right: [r.right - w, r.top, w, r.height] }[side];
        out.push({ t: "rect", r: rr, fill: c });
      }
      const before = getComputedStyle(el, "::before");
      if (/counter/.test(before.content)) {
        const reset = /\S+\s+(\d+)/.exec(getComputedStyle(el.parentElement).counterReset);
        const n = [...el.parentElement.children].indexOf(el) + 1 + (reset ? +reset[1] : 0);
        out.push({ t: "text", svg: true, align: "left",
          r: [r.left + (parseFloat(before.left) || 0), r.top + (parseFloat(before.top) || 0), 60,
              parseFloat(before.fontSize) * 1.3],
          runs: [{ text: String(n), size: parseFloat(before.fontSize),
                   bold: parseInt(before.fontWeight) >= 600, color: rgb(before.color) }] });
      }
    }
    if (/flex|grid/.test(s.display)) {
      for (const n of el.childNodes) {
        if (n.nodeType !== 3 || !n.textContent.trim()) continue;
        const range = document.createRange(); range.selectNodeContents(n);
        out.push({ t: "text", svg: true, align: "left", r: box(range.getBoundingClientRect()),
                   runs: [{ text: n.textContent.trim().replace(/\s+/g, " "), ...style(el) }] });
      }
    } else if (!inline(el) && el.tagName !== "IMG") {
      const runs = [], rects = [];
      collect(el, runs, rects);
      if (runs.some(x => x.text.trim()) && rects.length) {
        const pl = parseFloat(s.paddingLeft) + parseFloat(s.borderLeftWidth);
        const pr = parseFloat(s.paddingRight) + parseFloat(s.borderRightWidth);
        const lh = parseFloat(s.lineHeight) || parseFloat(s.fontSize) * 1.2;
        const top = Math.min(...rects.map(q => q.top)), bottom = Math.max(...rects.map(q => q.bottom));
        const lead = (lh - rects[0].height) / 2;
        out.push({ t: "text", r: [r.left + pl, top - lead, r.width - pl - pr, bottom - top + 2 * lead],
                   runs, lh, align: s.textAlign });
      }
    }
    for (const c of el.children) walk(c);
  }
  walk(document.documentElement);
  const tag = document.createElement("script");
  tag.type = "application/json"; tag.id = "probe"; tag.textContent = JSON.stringify(out);
  document.body.appendChild(tag);
});
</script>
"""


def probe(name: str) -> list[dict]:
    """Lay the body page out in Chrome and read back every element's box and style."""
    html = (FIGURES / f"{name}.html").read_text(encoding="utf-8")
    tmp = FIGURES / f"_probe-{name}.html"   # beside deck.css, so the page styles resolve
    tmp.write_text(html.replace("</head>", PROBE + "</head>"), encoding="utf-8")
    try:
        dom = subprocess.run(
            [str(CHROME), "--headless=new", "--disable-gpu", "--hide-scrollbars",
             "--window-size=4000,1710", "--virtual-time-budget=10000", "--dump-dom", str(tmp)],
            check=True, capture_output=True, text=True, encoding="utf-8").stdout
    finally:
        tmp.unlink(missing_ok=True)
    found = re.search(r'<script type="application/json" id="probe">(.*?)</script>', dom, re.S)
    if not found:
        raise SystemExit(f"{name}: the probe wrote nothing; is Chrome at {CHROME}?")
    return json.loads(found.group(1))


# ------------------------------------------------------------------ placing shapes

def at(x, y):
    return Inches(x / DPI), Inches(BODY_TOP + y / DPI)


def size(w, h):
    return Inches(max(w, 1) / DPI), Inches(max(h, 1) / DPI)


def colour(c):
    return RGBColor(*(round(v) for v in c))


def rect(slide, item, shape=MSO_SHAPE.RECTANGLE):
    x, y, w, h = item["r"]
    s = slide.shapes.add_shape(shape, *at(x, y), *size(w, h))
    if item.get("fill"):
        s.fill.solid()
        s.fill.fore_color.rgb = colour(item["fill"])
    else:
        s.fill.background()
    if item.get("stroke"):
        s.line.color.rgb = colour(item["stroke"])
        s.line.width = Pt(item["sw"] * PT)
    else:
        s.line.fill.background()
    s.shadow.inherit = False


def line(slide, item):
    pts = item["pts"]
    for i, ((x0, y0), (x1, y1)) in enumerate(zip(pts, pts[1:])):
        c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, *at(x0, y0), *at(x1, y1))
        c.line.color.rgb = colour(item["color"] or (79, 91, 105))
        c.line.width = Pt(item["sw"] * PT)
        if item["arrow"] and i == len(pts) - 2:
            ln = c.line._get_or_add_ln()
            ln.append(ln.makeelement(qn("a:tailEnd"), {"type": "triangle", "w": "med", "len": "med"}))


def picture(slide, item):
    x, y, w, h = item["r"]
    pic = slide.shapes.add_picture(str(FIGURES / item["src"]), *at(x, y), *size(w, h))
    if item["fit"] != "cover":
        return
    nw, nh = item["nw"], item["nh"]
    scale = max(w / nw, h / nh)
    vw, vh = w / scale, h / scale
    px, py = (float(v.rstrip("%")) / 100 if v.endswith("%") else 0.5 for v in item["pos"].split())
    left, top = (nw - vw) * px, (nh - vh) * py
    pic.crop_left, pic.crop_right = left / nw, 1 - (left + vw) / nw
    pic.crop_top, pic.crop_bottom = top / nh, 1 - (top + vh) / nh


ALIGN = {"center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT, "end": PP_ALIGN.RIGHT}


def text(slide, item):
    x, y, w, h = item["r"]
    svg = item.get("svg")
    # Slack on width only: PowerPoint's Arial and Chrome's differ by a pixel here and there,
    # and a box exactly as wide as its longest line can wrap one word early.
    slack = 40 if svg else w * 0.03
    if svg and item["align"] == "center":
        x -= slack / 2
    elif svg and item["align"] == "right":
        x -= slack
    box = slide.shapes.add_textbox(*at(x, y), *size(w + slack, h))
    frame = box.text_frame
    frame.word_wrap = not svg
    frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
    para = frame.paragraphs[0]
    para.alignment = ALIGN.get(item["align"], PP_ALIGN.LEFT)
    if item.get("lh"):
        para.line_spacing = Pt(item["lh"] * PT)
    runs = item["runs"]
    runs[0]["text"] = runs[0]["text"].lstrip()
    runs[-1]["text"] = runs[-1]["text"].rstrip()
    for spec in runs:
        if not spec["text"]:
            continue
        if spec["text"] == "\n":
            first = frame.paragraphs[0]
            para = frame.add_paragraph()
            para.alignment, para.line_spacing = first.alignment, first.line_spacing
            continue
        run = para.add_run()
        run.text = spec["text"].upper() if spec.get("upper") else spec["text"]
        font = run.font
        font.name = MONO if spec.get("mono") else FONT
        rPr = run._r.get_or_add_rPr()
        for slot in ("a:ea", "a:cs"):
            rPr.append(rPr.makeelement(qn(slot), {"typeface": font.name}))
        font.size = Pt(round(spec["size"] * PT * 2) / 2)
        font.bold = spec.get("bold", False)
        font.italic = spec.get("italic", False)
        if spec.get("color"):
            font.color.rgb = colour(spec["color"])


PLACE = {"rect": rect, "diamond": lambda s, i: rect(s, i, MSO_SHAPE.DIAMOND),
         "line": line, "img": picture, "text": text}


def main() -> None:
    deck = Presentation(str(TEMPLATE))
    fill_title(deck.slides[0])
    fill_idea_title(deck.slides[1])
    for slide in deck.slides:
        place_logo(slide)

    for index, name in BODIES.items():
        slide = deck.slides[index]
        strip_body(slide)
        items = probe(name)
        for item in items:
            PLACE[item["t"]](slide, item)
        print(f"{name}: {len(items)} shapes")

    xml_slides = deck.slides._sldIdLst
    slides = list(xml_slides)
    for index in sorted(DROP, reverse=True):
        xml_slides.remove(slides[index])

    OUT.parent.mkdir(parents=True, exist_ok=True)
    deck.save(str(OUT))
    print(f"wrote {OUT}")


if __name__ == "__main__":
    main()
