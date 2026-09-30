"""finalv2: finalv1 plus WCS on the roadmap, the one brief item the deck did not mention."""
from pptx import Presentation

SRC = "submission/sih/final/VVater-SIH2026-finalv1.pptx"
OUT = "submission/sih/final/VVater-SIH2026-finalv2.pptx"
p = Presentation(SRC)
six_mo = next(sh for sh in p.slides[3].shapes if sh.shape_id == 116)
run = six_mo.text_frame.paragraphs[1].runs
assert len(run) == 1 and run[0].text == "INCOIS servers, SSO", [r.text for r in run]
run[0].text = "INCOIS servers, SSO, OGC WCS"
p.save(OUT)
print("wrote", OUT)
